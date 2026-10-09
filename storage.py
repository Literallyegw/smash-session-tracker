# storage.py

import os
import json
from datetime import date
from openpyxl import Workbook, load_workbook
from constants import CONFIG_FILE, EXCEL_FILE, GAME_DATA, BASE_HEADERS, get_sheet_headers

DEFAULT_CONFIG = {
    "user_tag": "Player",
    "last_game": "SSBU",
    "last_char_ssbu": "",
    "last_char_ssbm": "",
    "known_opponents": []
}

def load_config() -> dict:
    if not os.path.exists(CONFIG_FILE):
        save_config(DEFAULT_CONFIG)
        return DEFAULT_CONFIG.copy()
    try:
        with open(CONFIG_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
            for k, v in DEFAULT_CONFIG.items():
                data.setdefault(k, v)
            return data
    except (json.JSONDecodeError, OSError):
        return DEFAULT_CONFIG.copy()

def save_config(config_data: dict) -> None:
    try:
        with open(CONFIG_FILE, "w", encoding="utf-8") as f:
            json.dump(config_data, f, indent=4, ensure_ascii=False)
    except OSError as e:
        print(f"Warning: Failed to save config: {e}")

def update_preferences(game: str = None, user_tag: str = None, my_char: str = None, new_opponent: str = None) -> dict:
    cfg = load_config()
    if game:
        cfg["last_game"] = game
    if user_tag is not None:
        cfg["user_tag"] = user_tag.strip()
    if my_char and game:
        if game == "SSBU":
            cfg["last_char_ssbu"] = my_char
        elif game == "SSBM":
            cfg["last_char_ssbm"] = my_char
    if new_opponent:
        opp_clean = new_opponent.strip()
        if opp_clean and opp_clean not in cfg.get("known_opponents", []):
            cfg.setdefault("known_opponents", []).append(opp_clean)
            cfg["known_opponents"].sort()

    save_config(cfg)
    return cfg

def init_excel_file() -> None:
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws_u = wb.active
        ws_u.title = "SSBU"
        ws_u.append(get_sheet_headers("SSBU"))

        ws_m = wb.create_sheet(title="SSBM")
        ws_m.append(get_sheet_headers("SSBM"))

        wb.save(EXCEL_FILE)

def save_session(game: str, match_type: str, my_tag: str, my_char: str,
                 opp_tag: str, opp_char: str, games_won: int, games_lost: int,
                 mental: int, lost_stages_dict: dict) -> str:
    init_excel_file()

    try:
        wb = load_workbook(EXCEL_FILE)
    except PermissionError:
        raise PermissionError("Το αρχείο Excel είναι ανοιχτό. Κλείσε το πρόγραμμα Excel και προσπάθησε ξανά.")

    if game not in wb.sheetnames:
        ws = wb.create_sheet(title=game)
        ws.append(get_sheet_headers(game))
    else:
        ws = wb[game]

    today_str = date.today().strftime("%Y-%m-%d")
    total_games = games_won + games_lost

    row = [
        match_type,
        today_str,
        my_tag,
        my_char,
        opp_tag,
        opp_char,
        total_games,
        games_won,
        games_lost,
        mental,
    ]

    for stage in GAME_DATA[game]["stages"]:
        row.append(1 if lost_stages_dict.get(stage, False) else 0)

    ws.append(row)

    try:
        wb.save(EXCEL_FILE)
    except PermissionError:
        raise PermissionError("Αδυναμία αποθήκευσης: Το Excel είναι ανοιχτό.")

    return today_str

def load_sessions(game: str) -> list[dict]:
    if not os.path.exists(EXCEL_FILE):
        return []

    try:
        wb = load_workbook(EXCEL_FILE, data_only=True)
    except PermissionError:
        raise PermissionError("Το αρχείο Excel είναι κλειδωμένο. Κλείσε το Excel.")
    except Exception:
        return []

    if game not in wb.sheetnames:
        return []

    ws = wb[game]
    records = []
    stages_list = GAME_DATA[game]["stages"]

    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(cell is not None and str(cell).strip() != "" for cell in row):
            lost_stages = []
            stage_flags = {}
            for idx, stage in enumerate(stages_list):
                col_idx = len(BASE_HEADERS) + idx
                is_lost = (col_idx < len(row) and row[col_idx] == 1)
                stage_flags[stage] = is_lost
                if is_lost:
                    lost_stages.append(stage)

            stages_str = ", ".join(lost_stages) if lost_stages else "None"

            records.append({
                "MatchType": str(row[0]) if len(row) > 0 and row[0] is not None else "",
                "Date": str(row[1]) if len(row) > 1 and row[1] is not None else "",
                "MyTag": str(row[2]) if len(row) > 2 and row[2] is not None else "",
                "MyChar": str(row[3]) if len(row) > 3 and row[3] is not None else "",
                "Opponent": str(row[4]) if len(row) > 4 and row[4] is not None else "",
                "Character": str(row[5]) if len(row) > 5 and row[5] is not None else "",
                "TotalGames": int(row[6]) if len(row) > 6 and row[6] is not None else 0,
                "Won": int(row[7]) if len(row) > 7 and row[7] is not None else 0,
                "Lost": int(row[8]) if len(row) > 8 and row[8] is not None else 0,
                "Mental": int(row[9]) if len(row) > 9 and row[9] is not None else 0,
                "StagesLost": stages_str,
                "StageFlags": stage_flags  # raw dict {stage_name: True/False}
            })

    return records

def delete_last_session(game: str) -> bool:
    if not os.path.exists(EXCEL_FILE):
        return False

    try:
        wb = load_workbook(EXCEL_FILE)
    except PermissionError:
        raise PermissionError("Το αρχείο Excel είναι ανοιχτό. Κλείσε το για να γίνει η διαγραφή.")

    if game not in wb.sheetnames:
        return False

    ws = wb[game]
    if ws.max_row <= 1:
        return False

    ws.delete_rows(ws.max_row, 1)

    try:
        wb.save(EXCEL_FILE)
    except PermissionError:
        raise PermissionError("Αδυναμία αποθήκευσης μετά τη διαγραφή: Το Excel είναι ανοιχτό.")

    return True